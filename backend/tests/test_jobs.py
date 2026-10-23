import io
import threading

from openpyxl import Workbook, load_workbook

from app.core.db import SessionLocal
from app.modules.jobs.canonical import canonicalize_url, dedupe_key, detect_ats, employer_key
from app.modules.jobs.models import JobSource
from app.modules.jobs.service import upsert_job


def test_canonicalization_and_ats_detection():
    a = canonicalize_url("HTTP://Boards.Greenhouse.io/acme/jobs/12345/?utm_source=x&gh_src=y#apply")
    assert a == "https://boards.greenhouse.io/acme/jobs/12345"
    ats = detect_ats(a)
    assert (ats.provider, ats.company, ats.requisition_id) == ("greenhouse", "acme", "12345")
    lever = detect_ats(
        canonicalize_url(
            "https://jobs.lever.co/Acme/0b1c2d3e-4f50-6172-8394-a5b6c7d8e9f0/apply?lever-source=li"
        )
    )
    assert lever.provider == "lever" and lever.company == "acme"
    embedded = detect_ats(canonicalize_url("https://careers.acme.example/jobs?gh_jid=987"))
    assert embedded.provider == "greenhouse" and embedded.requisition_id == "987"
    assert detect_ats("https://example.com/careers/1").provider == "unknown"


def test_employer_key_strips_legal_suffixes():
    assert employer_key("Acme, Inc.") == employer_key("ACME Inc") == employer_key("Acme") == "acme"


def test_dedupe_key_prefers_requisition():
    k1 = dedupe_key(
        canonical_url="https://boards.greenhouse.io/acme/jobs/1",
        employer="Acme",
        title="A",
        location=None,
    )
    k2 = dedupe_key(
        canonical_url="https://boards.greenhouse.io/acme/jobs/1",
        employer="Other",
        title="B",
        location="X",
    )
    assert k1 == k2
    assert dedupe_key(
        canonical_url=None, employer="Acme Inc", title="Nurse", location="Austin"
    ) == dedupe_key(canonical_url=None, employer="acme", title="  nurse ", location="AUSTIN")


def test_manual_job_and_duplicate(api):
    r = api.post(
        "/jobs",
        json={
            "url": "https://boards.greenhouse.io/acme/jobs/111?utm_source=li",
            "title": "Electrician",
            "employer": "Acme",
        },
    )
    assert r.status_code == 201 and r.json()["duplicate"] is False
    job = r.json()["job"]
    assert job["ats_provider"] == "greenhouse"
    r = api.post("/jobs", json={"url": "https://boards.greenhouse.io/acme/jobs/111"})
    assert r.status_code == 200 and r.json()["duplicate"] is True
    assert r.json()["job"]["id"] == job["id"]
    assert r.json()["job"]["seen_count"] == 2
    assert api.get("/jobs").json()["total"] == 1


def test_job_requires_url_or_title_and_employer(api):
    assert api.post("/jobs", json={"title": "Nurse"}).status_code == 422
    assert api.post("/jobs", json={"title": "Nurse", "employer": "Clinic"}).status_code == 201


def test_unsafe_job_urls_rejected(api):
    for url in (
        "http://127.0.0.1/admin",
        "http://169.254.169.254/latest/meta-data",
        "file:///etc/passwd",
        "http://localhost:8000",
        "http://10.0.0.5/",
        "http://user:pass@example.com/",
        "javascript:alert(1)",
        "http://[::1]/",
        "http://metadata.google.internal/",
    ):
        r = api.post("/jobs", json={"url": url})
        assert r.status_code == 400, url


def test_excluded_employer_gets_reason(api):
    api.put("/profile/preferences", json={"excluded_employers": ["Acme Inc."]})
    job = api.post("/jobs", json={"title": "Driver", "employer": "ACME"}).json()["job"]
    assert job["status"] == "excluded"
    assert "excluded list" in job["exclusion_reason"]
    detail = api.get(f"/jobs/{job['id']}").json()
    assert any(e["kind"] == "exclude" for e in detail["explanations"])


def test_match_explanations(api):
    api.put(
        "/profile/preferences",
        json={"titles": ["Teacher"], "target_locations": [{"city": "Denver", "country": "US"}]},
    )
    job = api.post(
        "/jobs",
        json={"title": "High School Teacher", "employer": "District 1", "location": "Denver, CO"},
    ).json()["job"]
    kinds = [e["kind"] for e in api.get(f"/jobs/{job['id']}").json()["explanations"]]
    assert kinds == ["match", "match"]


CSV = (
    "Job URL,Company,Job Title,Location,Status,Date Posted,Notes,Row ID\n"
    "https://jobs.lever.co/acme/0b1c2d3e-4f50-6172-8394-a5b6c7d8e9f0,Acme,Welder,Austin,,2026-10-01,,r1\n"
    "https://jobs.lever.co/acme/0b1c2d3e-4f50-6172-8394-a5b6c7d8e9f0?lever-source=x,Acme,Welder,Austin,,,,r2\n"
    ",Clinic,Nurse,Remote,shortlisted,,keen,r3\n"
    ",,,,,,,r4\n"
    ",,,,,,,\n"
    ",Clinic,,,,,,r5\n"
    "http://127.0.0.1/x,Bad,Bad,,,,,r6\n"
    ",Shop,Clerk,,bogus,10/01/2026,,r7\n"
)


def test_csv_import_row_outcomes(api):
    r = api.post("/jobs/import", files={"file": ("jobs.csv", CSV.encode(), "text/csv")})
    assert r.status_code == 200, r.text
    body = r.json()
    outcomes = {row["row_id"]: row["outcome"] for row in body["rows"]}
    assert outcomes == {
        "r1": "created",
        "r2": "duplicate",
        "r3": "created",
        "r4": "error",
        "r5": "error",
        "r6": "error",
        "r7": "error",
    }
    r7 = next(x for x in body["rows"] if x["row_id"] == "r7")
    assert len(r7["errors"]) == 2
    # Re-importing the same file creates nothing new.
    again = api.post("/jobs/import", files={"file": ("jobs.csv", CSV.encode(), "text/csv")}).json()
    assert again["created"] == 0 and again["duplicates"] == 3
    assert api.get("/jobs").json()["total"] == 2


def test_xlsx_import_rejects_formula_cells(api):
    wb = Workbook()
    ws = wb.active
    ws.append(["title", "employer", "notes"])
    ws.append(["Plumber", "Pipes Ltd", "fine"])
    ws.append(["Painter", "Brush Co", '=HYPERLINK("http://evil.example","x")'])
    buf = io.BytesIO()
    wb.save(buf)
    r = api.post(
        "/jobs/import", files={"file": ("jobs.xlsx", buf.getvalue(), "application/octet-stream")}
    )
    body = r.json()
    assert body["created"] == 1 and body["errors"] == 1
    assert "formula" in body["rows"][1]["errors"][0]


def test_import_rejects_bad_files(api):
    assert (
        api.post("/jobs/import", files={"file": ("x.txt", b"a", "text/plain")}).status_code == 400
    )
    r = api.post("/jobs/import", files={"file": ("x.csv", b"foo,bar\n1,2\n", "text/csv")})
    assert r.status_code == 400 and "recognised" in r.json()["detail"]


def test_export_neutralizes_formulas(api):
    api.post("/jobs", json={"title": "=cmd|' /C calc'!A0", "employer": "@SUM(1)", "notes": "+1-2"})
    csv_out = api.get("/jobs/export?format=csv").text
    assert "'=cmd" in csv_out and "'@SUM" in csv_out and "'+1-2" in csv_out
    xlsx = api.get("/jobs/export?format=xlsx").content
    ws = load_workbook(io.BytesIO(xlsx)).active
    values = [c.value for c in ws[2]]
    assert "'=cmd|' /C calc'!A0" in values
    assert all(c.data_type != "f" for c in ws[2])


def test_concurrent_upserts_create_one_job(api):
    user_id = api.get("/auth/me").json()["user_id"]
    results = []

    def worker():
        with SessionLocal() as db:
            _, created = upsert_job(
                db,
                user_id,
                source=JobSource.manual,
                url="https://jobs.ashbyhq.com/acme/0b1c2d3e-4f50-6172-8394-a5b6c7d8e9f0",
            )
            db.commit()
            results.append(created)

    threads = [threading.Thread(target=worker) for _ in range(8)]
    [t.start() for t in threads]
    [t.join() for t in threads]
    assert sorted(results) == [False] * 7 + [True]
    assert api.get("/jobs").json()["total"] == 1


def test_cross_tenant_job_access_denied(api, other_api):
    job = api.post("/jobs", json={"title": "Nurse", "employer": "Clinic"}).json()["job"]
    assert other_api.get(f"/jobs/{job['id']}").status_code == 404
    assert other_api.patch(f"/jobs/{job['id']}", json={"status": "closed"}).status_code == 404
    assert other_api.get("/jobs").json()["total"] == 0
    # Same posting imported by another user is their own, separate record.
    assert other_api.post("/jobs", json={"title": "Nurse", "employer": "Clinic"}).status_code == 201


def test_an_over_long_cell_is_that_rows_error_never_a_server_error(api):
    """Each stored column's limit (the job API's); a longer cell was a database error that
    failed the whole import (HTTP 500). Synthetic rows."""
    long = {"employer": "E" * 201, "title": "T" * 301, "location": "L" * 301,
            "requisition_id": "R" * 201}  # fmt: skip
    rows = ["employer,title,location,requisition_id,row_id"]
    rows.append("Acme Example (fictional),Clerk,,REQ-1,ok")
    for col in long:
        cells = {"employer": "Acme Example (fictional)", "title": "Clerk", "location": "",
                 "requisition_id": "REQ-2", col: long[col]}  # fmt: skip
        rows.append(",".join([cells["employer"], cells["title"], cells["location"],
                              cells["requisition_id"], col]))  # fmt: skip
    r = api.post("/jobs/import", files={"file": ("jobs.csv", "\n".join(rows).encode(), "text/csv")})
    assert r.status_code == 200, r.text
    out = {row["row_id"]: row for row in r.json()["rows"]}
    assert out["ok"]["outcome"] == "created"
    for col in long:
        assert out[col]["outcome"] == "error", col
        assert any(e.startswith(f"{col}: longer than") for e in out[col]["errors"]), out[col]
