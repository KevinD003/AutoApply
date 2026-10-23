"""CSV/XLSX import and export for jobs.

Formulas are never evaluated: XLSX formula cells are rejected per row, and exported cells that
a spreadsheet would interpret as formulas are prefixed with an apostrophe.
"""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

# canonical column -> accepted header aliases (case/space-insensitive)
COLUMNS: dict[str, set[str]] = {
    "row_id": {"row_id", "id", "row id"},
    "job_url": {"job_url", "url", "job url", "link", "posting url"},
    "employer": {"employer", "company", "company name", "organization"},
    "title": {"title", "job title", "position", "role"},
    "location": {"location", "city", "job location"},
    "status": {"status"},
    "posted_on": {"posted_on", "date posted", "posted", "posted date"},
    "requisition_id": {"requisition_id", "req id", "requisition", "job id"},
    "notes": {"notes", "note", "comments"},
}
EXPORT_COLUMNS = [
    "row_id",
    "job_url",
    "employer",
    "title",
    "location",
    "status",
    "posted_on",
    "requisition_id",
    "ats_provider",
    "first_seen_at",
    "notes",
]
ALLOWED_STATUSES = {"new", "shortlisted", "excluded", "closed"}
# The longest value each stored column takes (the same limits as the job API); a longer cell is
# that row's error, never a database error for the whole import.
MAX_LENGTHS = {"job_url": 2048, "employer": 200, "title": 300, "location": 300,
               "requisition_id": 200, "notes": 10_000, "row_id": 200}  # fmt: skip
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


class SpreadsheetError(ValueError):
    pass


@dataclass
class ParsedRow:
    row_number: int  # 1-based, header is row 1
    values: dict[str, str | None]
    errors: list[str] = field(default_factory=list)


def _norm_header(h: object) -> str:
    return " ".join(str(h or "").strip().lower().replace("_", " ").split())


def map_headers(headers: list[object]) -> dict[int, str]:
    lookup = {_norm_header(a): col for col, aliases in COLUMNS.items() for a in aliases}
    mapping: dict[int, str] = {}
    for i, h in enumerate(headers):
        col = lookup.get(_norm_header(h))
        if col and col not in mapping.values():
            mapping[i] = col
    if not ({"job_url", "title"} & set(mapping.values())):
        raise SpreadsheetError(
            "No recognised columns. Include at least 'job_url' or 'title' (see docs for the list)."
        )
    return mapping


def _cell_to_str(v: object) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def parse(content: bytes, file_format: str, max_rows: int) -> list[ParsedRow]:
    if file_format == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as e:
            raise SpreadsheetError("CSV must be UTF-8 encoded") from e
        table: list[list[object]] = list(csv.reader(io.StringIO(text)))
        formula_cells: set[tuple[int, int]] = set()
    elif file_format == "xlsx":
        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except Exception as e:  # openpyxl raises a variety of errors for bad files
            raise SpreadsheetError("Could not read XLSX file") from e
        ws = wb.worksheets[0]
        table, formula_cells = [], set()
        for r, row in enumerate(ws.iter_rows(), start=1):
            vals = []
            for c, cell in enumerate(row):
                if cell.data_type == "f":
                    formula_cells.add((r, c))
                    vals.append(None)
                else:
                    vals.append(cell.value)
            table.append(vals)
            if r > max_rows + 1:
                break
        wb.close()
    else:
        raise SpreadsheetError("Unsupported format; use .csv or .xlsx")

    if not table:
        raise SpreadsheetError("File is empty")
    mapping = map_headers(table[0])
    body = table[1:]
    if len(body) > max_rows:
        raise SpreadsheetError(f"Too many rows (max {max_rows})")

    out: list[ParsedRow] = []
    for idx, raw in enumerate(body, start=2):
        if all(_cell_to_str(v) is None for v in raw):
            continue
        values = {col: None for col in COLUMNS}
        errors: list[str] = []
        for i, col in mapping.items():
            if (idx, i) in formula_cells:
                errors.append(f"{col}: formula cells are not accepted")
                continue
            values[col] = _cell_to_str(raw[i]) if i < len(raw) else None
            if values[col] and col in MAX_LENGTHS and len(values[col]) > MAX_LENGTHS[col]:
                errors.append(f"{col}: longer than {MAX_LENGTHS[col]} characters")
                values[col] = None
        if values["status"]:
            values["status"] = values["status"].lower()
            if values["status"] not in ALLOWED_STATUSES:
                errors.append(f"status: must be one of {sorted(ALLOWED_STATUSES)}")
        if values["posted_on"]:
            try:
                values["posted_on"] = date.fromisoformat(values["posted_on"][:10]).isoformat()
            except ValueError:
                errors.append("posted_on: use YYYY-MM-DD")
        if not values["job_url"] and not (values["title"] and values["employer"]):
            errors.append("row needs job_url, or both title and employer")
        out.append(ParsedRow(idx, values, errors))
    return out


def neutralize(v: object) -> object:
    if isinstance(v, str) and v.startswith(_FORMULA_PREFIX):
        return "'" + v
    return v


def export_rows(rows: list[dict], file_format: str) -> bytes:
    if file_format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(EXPORT_COLUMNS)
        for r in rows:
            w.writerow(
                [neutralize(r.get(c)) if r.get(c) is not None else "" for c in EXPORT_COLUMNS]
            )
        return buf.getvalue().encode("utf-8")
    if file_format == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "jobs"
        ws.append(EXPORT_COLUMNS)
        for r in rows:
            ws.append([neutralize(r.get(c)) for c in EXPORT_COLUMNS])
        # Force every written cell to plain text so nothing is stored as a formula.
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    raise SpreadsheetError("Unsupported format; use csv or xlsx")
